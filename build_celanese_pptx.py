"""Build Celanese executive sample deck (illustrative)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUTPUT = r"C:\Users\neelk\Cursor\Celanese_Executive_Sample.pptx"
FONT = "Aptos"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN_L = Inches(0.72)
MARGIN_R = Inches(0.72)
MARGIN_T = Inches(0.55)
MARGIN_B = Inches(0.45)
CONTENT_W = SLIDE_W - MARGIN_L - MARGIN_R

ORANGE = RGBColor(0xE4, 0x50, 0x10)
TEAL = RGBColor(0x00, 0x78, 0xA8)
CHARCOAL = RGBColor(0x40, 0x40, 0x40)
BODY = RGBColor(0x1A, 0x1A, 0x1A)
SLATE = RGBColor(0x5C, 0x67, 0x70)
MIST = RGBColor(0xF4, 0xF5, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAD_START = RGBColor(0xFF, 0x8A, 0x1E)
GRAD_END = RGBColor(0x8C, 0x1A, 0x08)


def apply_gradient(fill, c_start, c_end, angle=135.0):
    fill.gradient()
    fill.gradient_angle = angle
    stops = fill.gradient_stops
    stops[0].position = 0.0
    stops[0].color.rgb = c_start
    stops[1].position = 1.0
    stops[1].color.rgb = c_end


def style_run(run, size_pt, color, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic


def add_paragraph(
    text_frame,
    text,
    size_pt,
    color,
    bold=False,
    italic=False,
    space_before=0,
    space_after=6,
    align=PP_ALIGN.LEFT,
):
    if text_frame.paragraphs[0].text == "" and len(text_frame.paragraphs) == 1:
        p = text_frame.paragraphs[0]
    else:
        p = text_frame.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    style_run(run, size_pt, color, bold=bold, italic=italic)
    return p


def text_box(slide, left, top, width, height):
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(0)
    return shape


def slide_background_gradient(slide, c_start, c_end, angle=135.0):
    fill = slide.background.fill
    apply_gradient(fill, c_start, c_end, angle=angle)


def add_oval_arc(slide, left, top, width, height, line_pt=1.25):
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, width, height)
    oval.fill.background()
    oval.line.color.rgb = WHITE
    oval.line.width = Pt(line_pt)
    return oval


def add_footer(slide, page_num):
    y = SLIDE_H - Inches(0.38)
    wm = text_box(slide, MARGIN_L, y, Inches(1.4), Inches(0.28))
    add_paragraph(wm.text_frame, "Celanese", 9, SLATE, bold=True)
    note = text_box(slide, Inches(3.2), y, Inches(7.5), Inches(0.28))
    add_paragraph(
        note.text_frame,
        "Illustrative figures. Not actual results.",
        8,
        SLATE,
        align=PP_ALIGN.CENTER,
    )
    pg = text_box(slide, SLIDE_W - MARGIN_R - Inches(0.55), y, Inches(0.55), Inches(0.28))
    add_paragraph(pg.text_frame, str(page_num), 9, SLATE, align=PP_ALIGN.RIGHT)


def build_cover(prs):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    slide_background_gradient(slide, GRAD_START, GRAD_END, angle=120.0)

    add_oval_arc(slide, Inches(-1.2), Inches(0.8), Inches(9.5), Inches(9.5), line_pt=1.0)
    add_oval_arc(slide, Inches(7.8), Inches(-2.0), Inches(11.0), Inches(11.0), line_pt=1.25)

    wm = text_box(slide, MARGIN_L, MARGIN_T, Inches(2.5), Inches(0.45))
    add_paragraph(wm.text_frame, "Celanese", 22, WHITE, bold=True)

    block_w = Inches(5.8)
    block_left = SLIDE_W - MARGIN_R - block_w
    block_top = SLIDE_H - Inches(2.85)
    title = text_box(slide, block_left, block_top, block_w, Inches(2.2))
    tf = title.text_frame
    add_paragraph(tf, "Specialty Polymers", 36, WHITE, bold=True, space_after=4)
    add_paragraph(tf, "Outlook", 36, WHITE, bold=True, space_after=14)
    add_paragraph(tf, "Illustrative executive update", 16, WHITE, space_after=6)
    add_paragraph(tf, "October 2026", 14, WHITE, bold=False)


def build_results_slide(prs):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    y = MARGIN_T
    eyebrow = text_box(slide, MARGIN_L, y, CONTENT_W, Inches(0.28))
    add_paragraph(eyebrow.text_frame, "ILLUSTRATIVE UPDATE", 10, TEAL, bold=True)
    y += Inches(0.32)

    title = text_box(slide, MARGIN_L, y, CONTENT_W, Inches(0.85))
    add_paragraph(
        title.text_frame,
        "Illustrative results point to a stronger cash start.",
        28,
        CHARCOAL,
        bold=True,
        space_after=0,
    )
    y += Inches(1.05)

    col_gap = Inches(0.45)
    left_w = Inches(5.35)
    right_w = CONTENT_W - left_w - col_gap
    left_x = MARGIN_L
    right_x = MARGIN_L + left_w + col_gap
    content_bottom = SLIDE_H - Inches(1.35)

    kpis = [
        ("$48 M", "Free cash flow"),
        ("$1.10", "Adjusted EPS"),
        ("$186 M", "EM adjusted EBIT"),
        ("$124 M", "AC adjusted EBIT"),
    ]
    kpi_h = Inches(1.05)
    kpi_gap = Inches(0.22)
    kpi_y = y
    for value, label in kpis:
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left_x, kpi_y, left_w, kpi_h
        )
        card.fill.solid()
        card.fill.fore_color.rgb = MIST
        card.line.fill.background()
        card.adjustments[0] = 0.08

        inner = text_box(slide, left_x + Inches(0.28), kpi_y + Inches(0.12), left_w - Inches(0.5), kpi_h - Inches(0.2))
        tf = inner.text_frame
        add_paragraph(tf, value, 26, ORANGE, bold=True, space_after=2)
        add_paragraph(tf, label, 13, BODY, space_after=2)
        add_paragraph(tf, "Illustrative", 9, SLATE, italic=True, space_after=0)
        kpi_y += kpi_h + kpi_gap

    highlights = text_box(slide, right_x, y, right_w, content_bottom - y)
    tf = highlights.text_frame
    add_paragraph(tf, "Highlights", 14, CHARCOAL, bold=True, space_after=10)
    bullets = [
        "Free cash flow improved versus the prior illustrative quarter on working-capital discipline.",
        "Mix and cost actions supported adjusted EPS despite flat illustrative volumes.",
        "Engineered Materials and Acetyl Chain both contributed to the illustrative EBIT uplift.",
    ]
    for i, bullet in enumerate(bullets):
        p = add_paragraph(tf, bullet, 13, BODY, space_after=10 if i < len(bullets) - 1 else 0)
        p.level = 0
        p.space_before = Pt(0)

    bar_h = Inches(0.62)
    bar_top = SLIDE_H - Inches(1.02)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), bar_top, SLIDE_W, bar_h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()

    bar_text = text_box(slide, MARGIN_L, bar_top + Inches(0.14), CONTENT_W, Inches(0.4))
    add_paragraph(
        bar_text.text_frame,
        "Illustrative takeaway: cash improved on mix and cost, with no change to the plan.",
        13,
        WHITE,
        bold=True,
    )

    add_footer(slide, 2)


def build_decision_slide(prs):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    y = MARGIN_T
    eyebrow = text_box(slide, MARGIN_L, y, CONTENT_W, Inches(0.28))
    add_paragraph(eyebrow.text_frame, "DECISION", 10, TEAL, bold=True)
    y += Inches(0.32)

    title = text_box(slide, MARGIN_L, y, CONTENT_W, Inches(0.75))
    add_paragraph(
        title.text_frame,
        "Proceed with the next cash and cost actions.",
        28,
        CHARCOAL,
        bold=True,
    )
    y += Inches(0.95)

    rec = text_box(slide, MARGIN_L, y, CONTENT_W, Inches(0.55))
    add_paragraph(
        rec.text_frame,
        "We recommend endorsing the illustrative plan to protect cash through year-end while completing targeted cost work.",
        14,
        BODY,
    )
    y += Inches(0.72)

    col_w = (CONTENT_W - Inches(0.4)) / 2
    left_x = MARGIN_L
    right_x = MARGIN_L + col_w + Inches(0.4)
    col_top = y
    col_h = Inches(2.55)

    for x, heading, items in (
        (
            left_x,
            "Do now",
            [
                "Close the sample cost program by 31 Oct 2026.",
                "Name a single cash owner for weekly illustrative tracking.",
            ],
        ),
        (
            right_x,
            "Hold",
            [
                "No site decision until the next illustrative review.",
                "No change to segment targets in this update.",
            ],
        ),
    ):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, col_top, col_w, col_h)
        box.fill.solid()
        box.fill.fore_color.rgb = MIST
        box.line.fill.background()
        box.adjustments[0] = 0.06
        inner = text_box(slide, x + Inches(0.28), col_top + Inches(0.22), col_w - Inches(0.56), col_h - Inches(0.35))
        tf = inner.text_frame
        add_paragraph(tf, heading, 16, TEAL, bold=True, space_after=12)
        for i, item in enumerate(items):
            p = add_paragraph(tf, item, 13, BODY, space_after=8 if i == 0 else 0)
            p.level = 0

    bar_h = Inches(0.62)
    bar_top = SLIDE_H - Inches(1.02)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), bar_top, SLIDE_W, bar_h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()

    ask = text_box(slide, MARGIN_L, bar_top + Inches(0.12), CONTENT_W, Inches(0.45))
    add_paragraph(
        ask.text_frame,
        "Decision requested: endorse the illustrative plan.  Owner: Finance.  Date: 31 Oct 2026.",
        12,
        WHITE,
        bold=True,
    )

    add_footer(slide, 3)


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    build_cover(prs)
    build_results_slide(prs)
    build_decision_slide(prs)
    prs.save(OUTPUT)
    verify = Presentation(OUTPUT)
    assert len(verify.slides) == 3
    print(f"Saved: {OUTPUT}")
    print(f"Slides: {len(verify.slides)}")
    print(f"Size: {verify.slide_width.inches:.3f}in x {verify.slide_height.inches:.3f}in")


if __name__ == "__main__":
    main()
