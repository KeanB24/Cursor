"""Generate a 15-min Site leadership debrief PPT from NASSCOM event notes."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# Brand palette — professional navy / teal (avoid purple AI cliché)
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
TEAL = RGBColor(0x0D, 0x7A, 0x7A)
ACCENT = RGBColor(0xE8, 0x6A, 0x17)  # warm amber accent
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF4, 0xF7, 0xFA)
MUTED = RGBColor(0x5A, 0x6A, 0x7A)
DARK = RGBColor(0x1A, 0x2A, 0x3A)
GREEN = RGBColor(0x1B, 0x7A, 0x4A)
RED_SOFT = RGBColor(0xB0, 0x3A, 0x2E)
CARD = RGBColor(0xE8, 0xEE, 0xF4)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _set_run_font(run, size=18, bold=False, color=DARK, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def _add_bg(slide, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def _bar(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def _textbox(slide, left, top, width, height, text, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run_font(run, size=size, bold=bold, color=color)
    return box


def _bullets(slide, left, top, width, height, items, size=16, color=DARK, bold_first=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        p.level = 0
        run = p.add_run()
        run.text = "•  " + item
        _set_run_font(run, size=size, bold=(bold_first and i == 0), color=color)
    return box


def _card(slide, left, top, width, height, fill=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(0xD0, 0xD8, 0xE0)
    shape.line.width = Pt(1)
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def _footer(slide, page, total=10):
    _textbox(
        slide,
        Inches(0.5),
        Inches(7.1),
        Inches(10),
        Inches(0.3),
        "NASSCOM Session Debrief  |  Confidential — Site Leadership",
        size=10,
        color=MUTED,
    )
    _textbox(
        slide,
        Inches(11.5),
        Inches(7.1),
        Inches(1.3),
        Inches(0.3),
        f"{page} / {total}",
        size=10,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    total = 10

    # ---- 1. Title ----
    s = prs.slides.add_slide(blank)
    _add_bg(s, NAVY)
    _bar(s, 0, Inches(0), Inches(0.18), SLIDE_H, TEAL)
    _textbox(s, Inches(0.8), Inches(1.6), Inches(11.5), Inches(0.4), "SITE LEADERSHIP DEBRIEF  ·  15 MINUTES", size=14, bold=True, color=TEAL)
    _textbox(
        s,
        Inches(0.8),
        Inches(2.1),
        Inches(11.5),
        Inches(1.6),
        "Human in the Loop Is Not Enough:\nWho Is Really Running the Workflow?",
        size=36,
        bold=True,
        color=WHITE,
    )
    _textbox(
        s,
        Inches(0.8),
        Inches(4.0),
        Inches(11.5),
        Inches(0.5),
        "NASSCOM session  ·  29 Sep 2026  ·  HICC, Hyderabad  ·  Ahead of NTC 2026",
        size=16,
        color=CARD,
    )
    _textbox(
        s,
        Inches(0.8),
        Inches(5.0),
        Inches(11.5),
        Inches(0.8),
        "Insights for how AI is reshaping enterprise workflows, accountability,\nand the skills / roles our site should prioritize through 2028.",
        size=16,
        color=CARD,
    )
    _textbox(s, Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.3), "Internal use — synthesized from session notes + public industry context", size=11, color=MUTED)

    # ---- 2. Agenda ----
    s = prs.slides.add_slide(blank)
    _add_bg(s, LIGHT)
    _bar(s, 0, 0, SLIDE_W, Inches(0.9), NAVY)
    _textbox(s, Inches(0.6), Inches(0.25), Inches(10), Inches(0.5), "Agenda — 15 minutes", size=26, bold=True, color=WHITE)
    agenda = [
        ("01", "Session snapshot & the central question", "~2 min"),
        ("02", "Why “human in the loop” is no longer enough", "~3 min"),
        ("03", "Who really runs the workflow in an agentic world", "~3 min"),
        ("04", "Roles accelerating vs. slowing (to 2028)", "~4 min"),
        ("05", "How we prepare — implications for the site", "~3 min"),
    ]
    for i, (num, title, mins) in enumerate(agenda):
        y = Inches(1.3) + Inches(i * 1.0)
        _card(s, Inches(0.6), y, Inches(12.1), Inches(0.85), WHITE)
        _textbox(s, Inches(0.9), y + Inches(0.2), Inches(1.0), Inches(0.45), num, size=22, bold=True, color=TEAL)
        _textbox(s, Inches(2.0), y + Inches(0.22), Inches(8.5), Inches(0.45), title, size=18, bold=True, color=DARK)
        _textbox(s, Inches(10.8), y + Inches(0.25), Inches(1.5), Inches(0.4), mins, size=14, color=MUTED, align=PP_ALIGN.RIGHT)
    _footer(s, 2, total)

    # ---- 3. Session snapshot ----
    s = prs.slides.add_slide(blank)
    _add_bg(s, LIGHT)
    _bar(s, 0, 0, SLIDE_W, Inches(0.9), NAVY)
    _textbox(s, Inches(0.6), Inches(0.25), Inches(12), Inches(0.5), "Session snapshot", size=26, bold=True, color=WHITE)

    _textbox(s, Inches(0.6), Inches(1.15), Inches(12), Inches(0.4), "Panelists", size=14, bold=True, color=TEAL)
    panelists = [
        ("Pradeep Pasupuleti", "Director, ML & Big Data — Hitachi Digital Services\n(also shared as Assistant Director, ML & Big Data, Hitachi India)"),
        ("Vijaya Kadiyala", "Exec. Director, Enterprise Data / AI / Cloud — DBS Tech India\n(also shared as Director of AI/ML Engineering)"),
        ("Sajo Mathews", "ML Lead / Director of AI — Deccan AI\nFocus: model evaluation, post-training, reliable real-world AI"),
    ]
    for i, (name, role) in enumerate(panelists):
        x = Inches(0.6) + Inches(i * 4.15)
        _card(s, x, Inches(1.55), Inches(3.95), Inches(2.35), WHITE)
        _bar(s, x, Inches(1.55), Inches(3.95), Inches(0.08), TEAL)
        _textbox(s, x + Inches(0.2), Inches(1.8), Inches(3.55), Inches(0.45), name, size=15, bold=True, color=NAVY)
        _textbox(s, x + Inches(0.2), Inches(2.35), Inches(3.55), Inches(1.3), role, size=12, color=MUTED)

    _card(s, Inches(0.6), Inches(4.15), Inches(12.1), Inches(2.5), WHITE)
    _textbox(s, Inches(0.9), Inches(4.35), Inches(11.5), Inches(0.35), "Central question the session explored", size=14, bold=True, color=TEAL)
    _textbox(
        s,
        Inches(0.9),
        Inches(4.8),
        Inches(11.5),
        Inches(1.5),
        "As AI systems take on greater autonomy, enterprise workflows are changing —\nfrom how decisions are made to how work gets done.\n\nIs a “human in the loop” still enough… or who is really accountable for outcomes?",
        size=16,
        color=DARK,
    )
    _footer(s, 3, total)

    # ---- 4. HITL not enough ----
    s = prs.slides.add_slide(blank)
    _add_bg(s, LIGHT)
    _bar(s, 0, 0, SLIDE_W, Inches(0.9), NAVY)
    _textbox(s, Inches(0.6), Inches(0.25), Inches(12), Inches(0.5), "Why “human in the loop” is not enough", size=26, bold=True, color=WHITE)

    _card(s, Inches(0.6), Inches(1.2), Inches(12.1), Inches(1.3), NAVY)
    _textbox(
        s,
        Inches(0.9),
        Inches(1.45),
        Inches(11.5),
        Inches(0.9),
        "Token approval checkpoints ≠ meaningful control.\nEnterprises buy systems, accountability, and outcomes — not just automated tasks.",
        size=18,
        bold=True,
        color=WHITE,
    )

    points = [
        ("From execution → ownership", "Humans move out of routine execution and into design, supervision, validation, and outcome ownership."),
        ("Accountability must be named", "Someone must own authority, escalation, and results — not merely click “Approve” under time pressure."),
        ("Regulated reality", "Legal liability, safety, compliance, and governance make fully autonomous end-states impractical for most enterprises."),
        ("HITL vs. HOTL", "Start with human-in-the-loop; graduate to human-on-the-loop only after proven bounds, monitoring, and evidence trails."),
    ]
    for i, (title, body) in enumerate(points):
        col = i % 2
        row = i // 2
        x = Inches(0.6) + Inches(col * 6.2)
        y = Inches(2.75) + Inches(row * 1.85)
        _card(s, x, y, Inches(5.95), Inches(1.7), WHITE)
        _textbox(s, x + Inches(0.25), y + Inches(0.25), Inches(5.4), Inches(0.4), title, size=15, bold=True, color=TEAL)
        _textbox(s, x + Inches(0.25), y + Inches(0.7), Inches(5.4), Inches(0.8), body, size=13, color=DARK)
    _footer(s, 4, total)

    # ---- 5. Who runs the workflow ----
    s = prs.slides.add_slide(blank)
    _add_bg(s, LIGHT)
    _bar(s, 0, 0, SLIDE_W, Inches(0.9), NAVY)
    _textbox(s, Inches(0.6), Inches(0.25), Inches(12), Inches(0.5), "Who is really running the workflow?", size=26, bold=True, color=WHITE)

    _textbox(
        s,
        Inches(0.6),
        Inches(1.15),
        Inches(12),
        Inches(0.45),
        "In an agentic operating model, “running the workflow” means owning the system — not doing every step.",
        size=15,
        color=MUTED,
    )

    layers = [
        ("Design", "Define goals, constraints,\npolicies, and success metrics"),
        ("Orchestration", "Integrate agents into legacy +\nmodern systems safely at scale"),
        ("Oversight", "Monitor exceptions, audit trails,\nescalate high-risk actions"),
        ("Accountability", "Named owner for outcomes,\ncompliance, and failure modes"),
    ]
    for i, (t, b) in enumerate(layers):
        x = Inches(0.5) + Inches(i * 3.2)
        _card(s, x, Inches(1.8), Inches(3.0), Inches(2.6), WHITE)
        _bar(s, x, Inches(1.8), Inches(3.0), Inches(0.1), ACCENT if i == 3 else TEAL)
        _textbox(s, x + Inches(0.2), Inches(2.15), Inches(2.6), Inches(0.45), f"{i+1}. {t}", size=16, bold=True, color=NAVY)
        _textbox(s, x + Inches(0.2), Inches(2.75), Inches(2.6), Inches(1.3), b, size=13, color=DARK)

    _card(s, Inches(0.5), Inches(4.7), Inches(12.3), Inches(1.9), WHITE)
    _textbox(s, Inches(0.8), Inches(4.9), Inches(11.8), Inches(0.35), "Leadership takeaway", size=14, bold=True, color=ACCENT)
    _textbox(
        s,
        Inches(0.8),
        Inches(5.35),
        Inches(11.8),
        Inches(1.0),
        "Process / product owners run the workflow. Agents execute tasks. Checkpoint approvers only matter if they have\nreal authority and enough information to intervene. High-risk actions may need dual control — legal responsibility\nstill sits with the organization and its accountable leaders.",
        size=14,
        color=DARK,
    )
    _footer(s, 5, total)

    # ---- 6. Roles growing ----
    s = prs.slides.add_slide(blank)
    _add_bg(s, LIGHT)
    _bar(s, 0, 0, SLIDE_W, Inches(0.9), NAVY)
    _textbox(s, Inches(0.6), Inches(0.25), Inches(12), Inches(0.5), "Roles expected to grow by 2028", size=26, bold=True, color=WHITE)
    _textbox(
        s,
        Inches(0.6),
        Inches(1.1),
        Inches(12),
        Inches(0.4),
        "From session notes — especially in the 2–5 year experience band",
        size=14,
        color=MUTED,
    )

    roles = [
        ("Forward Deployed Engineers (FDE)", "Embed with customers / domains; ship AI into real workflows; close the gap between model capability and business adoption."),
        ("AI Engineers", "Build, evaluate, and harden models & agent systems — data quality, evals, post-training, reliability in production."),
        ("Go-to-Market Engineers (GTM)", "Bridge product, demos, solutions, and sales engineering — translate AI capability into packaged value and adoption."),
    ]
    for i, (title, body) in enumerate(roles):
        y = Inches(1.6) + Inches(i * 1.55)
        _card(s, Inches(0.6), y, Inches(12.1), Inches(1.4), WHITE)
        _bar(s, Inches(0.6), y, Inches(0.12), Inches(1.4), GREEN)
        _textbox(s, Inches(1.0), y + Inches(0.25), Inches(11.2), Inches(0.4), title, size=18, bold=True, color=GREEN)
        _textbox(s, Inches(1.0), y + Inches(0.7), Inches(11.2), Inches(0.5), body, size=14, color=DARK)
    _footer(s, 6, total)

    # ---- 7. Roles slowing ----
    s = prs.slides.add_slide(blank)
    _add_bg(s, LIGHT)
    _bar(s, 0, 0, SLIDE_W, Inches(0.9), NAVY)
    _textbox(s, Inches(0.6), Inches(0.25), Inches(12), Inches(0.5), "Roles that may see slowdown", size=26, bold=True, color=WHITE)
    _textbox(
        s,
        Inches(0.6),
        Inches(1.1),
        Inches(12),
        Inches(0.4),
        "Not “disappearing overnight” — but demand and career leverage likely compress without reinvention",
        size=14,
        color=MUTED,
    )

    slow = [
        ("Entry-level / fresher roles", "Routine coding, documentation, and first-line analysis are increasingly AI-augmented. Fewer pure “learn by doing grunt work” seats."),
        ("Managers (PIPE) — limited build engagement", "People / process managers without direct development or domain-tech ownership risk becoming coordination overhead."),
        ("Production support (traditional)", "Ticket-driven, reactive L1/L2 work faces automation; value shifts to SRE-style ownership, reliability engineering, and incident judgment."),
    ]
    for i, (title, body) in enumerate(slow):
        y = Inches(1.6) + Inches(i * 1.55)
        _card(s, Inches(0.6), y, Inches(12.1), Inches(1.4), WHITE)
        _bar(s, Inches(0.6), y, Inches(0.12), Inches(1.4), RED_SOFT)
        _textbox(s, Inches(1.0), y + Inches(0.25), Inches(11.2), Inches(0.4), title, size=18, bold=True, color=RED_SOFT)
        _textbox(s, Inches(1.0), y + Inches(0.7), Inches(11.2), Inches(0.5), body, size=14, color=DARK)
    _footer(s, 7, total)

    # ---- 8. How to prepare ----
    s = prs.slides.add_slide(blank)
    _add_bg(s, LIGHT)
    _bar(s, 0, 0, SLIDE_W, Inches(0.9), NAVY)
    _textbox(s, Inches(0.6), Inches(0.25), Inches(12), Inches(0.5), "How to prepare for the shift", size=26, bold=True, color=WHITE)

    prep = [
        ("Tech (AI) + Domain", "Pair AI fluency with deep domain expertise in one’s role — that’s where durable leverage sits."),
        ("Thought leadership", "Move from task delivery to framing problems, writing, teaching, and shaping how the org uses AI."),
        ("Ideation & design", "Invest in problem framing, workflow redesign, and outcome definition — not just ticket completion."),
        ("Curiosity as habit", "Continuous experimentation with tools, evals, and agent patterns; stay uncomfortable on purpose."),
    ]
    for i, (t, b) in enumerate(prep):
        col = i % 2
        row = i // 2
        x = Inches(0.6) + Inches(col * 6.2)
        y = Inches(1.25) + Inches(row * 2.5)
        _card(s, x, y, Inches(5.95), Inches(2.25), WHITE)
        _bar(s, x, y, Inches(5.95), Inches(0.1), TEAL)
        _textbox(s, x + Inches(0.3), y + Inches(0.35), Inches(5.3), Inches(0.45), t, size=18, bold=True, color=NAVY)
        _textbox(s, x + Inches(0.3), y + Inches(1.0), Inches(5.3), Inches(1.0), b, size=15, color=DARK)
    _footer(s, 8, total)

    # ---- 9. Site implications ----
    s = prs.slides.add_slide(blank)
    _add_bg(s, LIGHT)
    _bar(s, 0, 0, SLIDE_W, Inches(0.9), NAVY)
    _textbox(s, Inches(0.6), Inches(0.25), Inches(12), Inches(0.5), "Implications for Site leadership", size=26, bold=True, color=WHITE)

    impl = [
        "Hire / grow for FDE, AI Engineering, and GTM-engineering profiles in the mid-band (2–5 yrs) — not only classic IC ladders.",
        "Redesign fresher & support pathways: AI-augmented onboarding, domain immersion, and measurable ownership early.",
        "Expect managers to stay technically / domain-engaged (build with teams) — pure PIPE coordination is a risk posture.",
        "For every agentic workflow: name an accountable owner, define escalation, and keep auditability before autonomy.",
        "Reward curiosity, ideation, and thought leadership in performance systems — not only utilization and ticket volume.",
        "Treat services value as system + outcome ownership (integration, risk, compliance) — not undifferentiated effort.",
    ]
    _bullets(s, Inches(0.7), Inches(1.3), Inches(12), Inches(5.3), impl, size=16, color=DARK)
    _footer(s, 9, total)

    # ---- 10. Close ----
    s = prs.slides.add_slide(blank)
    _add_bg(s, NAVY)
    _bar(s, 0, Inches(0), Inches(0.18), SLIDE_H, TEAL)
    _textbox(s, Inches(0.8), Inches(1.4), Inches(11.5), Inches(0.4), "CLOSING TAKEAWAY", size=14, bold=True, color=TEAL)
    _textbox(
        s,
        Inches(0.8),
        Inches(1.9),
        Inches(11.5),
        Inches(1.8),
        "AI will run more of the work.\nHumans must still own the outcome.",
        size=32,
        bold=True,
        color=WHITE,
    )
    _textbox(
        s,
        Inches(0.8),
        Inches(4.0),
        Inches(11.5),
        Inches(1.2),
        "Our competitive edge through 2028: people who combine AI capability with domain depth,\naccountable workflow ownership, and the curiosity to redesign how work gets done.",
        size=16,
        color=CARD,
    )
    _textbox(
        s,
        Inches(0.8),
        Inches(5.5),
        Inches(11.5),
        Inches(0.8),
        "Discussion: Where should we pilot FDE / AI-engineer pathways first?\nWhich workflows need a named accountable owner before we automate further?",
        size=15,
        color=ACCENT,
    )
    _textbox(s, Inches(0.8), Inches(6.7), Inches(11.5), Inches(0.3), "Sources: Session notes · NASSCOM NTC Hyderabad 2026 preview · NASSCOM Voices (Agentic Shift) · industry HITL/HOTL guidance", size=10, color=MUTED)

    out = Path(__file__).resolve().parent / "NASSCOM_HITL_Site_Leadership_Debrief.pptx"
    prs.save(out)
    print(f"Saved: {out}")
    return out


if __name__ == "__main__":
    build()
